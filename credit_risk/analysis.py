"""Predict whether a loan defaults within twelve months."""

# Libraries
from datetime import datetime
from pathlib import Path
from urllib.request import urlretrieve

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.ticker import PercentFormatter
from openpyxl import load_workbook
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import HistGradientBoostingClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import brier_score_loss, roc_auc_score
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler


# Files and analysis settings
ROOT = Path(__file__).resolve().parent
EXTRAS_DIR = ROOT.parent / "extras"
SOURCE = ROOT / "loan_dataset_investor.xlsx"
RESULTS = ROOT / "credit_results.xlsx"
FIGURE = EXTRAS_DIR / "credit_figure.pdf"
DATA_URL = (
    "https://sabanners001.blob.core.windows.net/statistics/public/"
    "loan_dataset_investor.xlsx"
)

START_DATE = datetime(2023, 10, 1)
VALIDATION_DATE = datetime(2024, 7, 1)
TEST_DATE = datetime(2024, 10, 1)
END_DATE = datetime(2025, 4, 1)
COUNTRIES = {"Estonia", "Finland", "Latvia", "Netherlands", "Denmark"}

SOURCE_COLUMNS = (
    "loan_id",
    "country",
    "loan_issued_at",
    "issued_amount",
    "initial_interest_rate",
    "initial_loan_duration",
    "combined_income",
    "has_default_within_12_months",
    "projected_npv_return",
    "customer_risk_rating",
)
RATING_ORDER = {
    grade: number
    for number, grade in enumerate(["AA", "A", "B", "C", "D", "E", "F", "G", "HR"])
}
MODEL_NAMES = {
    "public_logistic": "Logistic regression with public features",
    "public_pricing_logistic": "Logistic regression with public and pricing features",
    "rating_logistic": "Logistic regression with risk rating",
    "public_pricing_boosting": "Gradient boosting with public and pricing features",
}
PERIOD_NAMES = {
    "validation": "Validation set",
    "test": "Test set",
}


# Step 1: Read the original loan data

def download_data():
    if not SOURCE.exists():
        urlretrieve(DATA_URL, SOURCE)


def load_data():
    download_data()
    workbook = load_workbook(SOURCE, read_only=True)
    sheet = workbook["Loan Dataset"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    positions = {name: header.index(name) for name in SOURCE_COLUMNS}
    records = []

    for row in rows:
        values = {name: row[position] for name, position in positions.items()}
        if any(value is None for value in values.values()):
            continue

        issue_date = pd.to_datetime(values["loan_issued_at"]).to_pydatetime()
        if not START_DATE <= issue_date < END_DATE:
            continue
        if values["country"] not in COUNTRIES:
            continue

        values["loan_issued_at"] = issue_date
        records.append(values)
    workbook.close()

    loans = pd.DataFrame.from_records(records, columns=SOURCE_COLUMNS)
    loans = loans.rename(
        columns={"has_default_within_12_months": "default_within_12_months"}
    )
    loans["default_within_12_months"] = loans[
        "default_within_12_months"
    ].astype(int)
    loans = loans.sort_values(["loan_issued_at", "loan_id"], ignore_index=True)
    loans["split"] = np.select(
        [
            loans["loan_issued_at"] < VALIDATION_DATE,
            loans["loan_issued_at"] < TEST_DATE,
        ],
        ["train", "validation"],
        default="test",
    )
    loans["loan_group"] = np.where(
        loans["country"].eq("Denmark"), "Denmark", "main"
    )
    return loans


# Step 2: Create the model features

def create_features(loans):
    loans = loans.copy()
    loans["log_issued_amount"] = np.log1p(loans["issued_amount"])
    loans["log_combined_income"] = np.log1p(loans["combined_income"])
    loans["rating_score"] = loans["customer_risk_rating"].map(RATING_ORDER)
    return loans


# Step 3: Define the four classifiers

def logistic_model(numeric, categorical):
    preprocessing = ColumnTransformer(
        [
            ("numeric", StandardScaler(), numeric),
            ("categorical", OneHotEncoder(handle_unknown="ignore"), categorical),
        ]
    )
    return Pipeline(
        [
            ("preprocessing", preprocessing),
            ("logistic regression", LogisticRegression()),
        ]
    )


def boosting_model():
    numeric = [
        "log_issued_amount",
        "initial_loan_duration",
        "log_combined_income",
        "initial_interest_rate",
        "projected_npv_return",
    ]
    preprocessing = ColumnTransformer(
        [
            ("numeric", "passthrough", numeric),
            (
                "country",
                OneHotEncoder(handle_unknown="ignore", sparse_output=False),
                ["country"],
            ),
        ]
    )
    model = HistGradientBoostingClassifier(
        learning_rate=0.05,
        max_iter=250,
        max_leaf_nodes=15,
        min_samples_leaf=100,
        l2_regularization=1.0,
        early_stopping=False,
    )
    return Pipeline([("preprocessing", preprocessing), ("gradient boosting", model)])


def make_models():
    public = ["log_issued_amount", "initial_loan_duration", "log_combined_income"]
    public_and_pricing = public + ["initial_interest_rate", "projected_npv_return"]
    return {
        "public_logistic": logistic_model(public, ["country"]),
        "public_pricing_logistic": logistic_model(public_and_pricing, ["country"]),
        "rating_logistic": logistic_model(["rating_score"], []),
        "public_pricing_boosting": boosting_model(),
    }


# Step 4: Train on earlier loans and evaluate later loans

def calculate_metrics(period, model_name, target, probability):
    return {
        "period": period,
        "model": model_name,
        "loans": len(target),
        "defaults": int(target.sum()),
        "roc_auc": roc_auc_score(target, probability),
        "brier_score": brier_score_loss(target, probability),
        "mean_probability": probability.mean(),
        "default_rate": target.mean(),
    }


def fit_models(train, validation, test):
    target = "default_within_12_months"
    models = make_models()
    metrics = []
    test_predictions = test[["loan_id", "loan_issued_at", "country", target]].copy()

    for name, model in models.items():
        model.fit(train, train[target])
        for period_name, data in (("validation", validation), ("test", test)):
            probability = model.predict_proba(data)[:, 1]
            metrics.append(
                calculate_metrics(period_name, name, data[target].to_numpy(), probability)
            )
            if period_name == "test":
                test_predictions[name] = probability

    return models, pd.DataFrame(metrics), test_predictions


# Step 5: Examine the selected logistic-regression probabilities

def calibration_table(test_predictions, denmark_predictions):
    probability = "public_pricing_logistic"
    target = "default_within_12_months"
    frames = []

    risk_group = pd.qcut(test_predictions[probability], 10, labels=False) + 1
    groupings = [
        ("Risk group", test_predictions.assign(group=risk_group)),
        ("Country", test_predictions.assign(group=test_predictions["country"])),
        (
            "Issue quarter",
            test_predictions.assign(
                group=pd.to_datetime(test_predictions["loan_issued_at"])
                .dt.to_period("Q")
                .astype(str)
            ),
        ),
        ("Country outside training", denmark_predictions.assign(group="Denmark")),
    ]

    for check_name, data in groupings:
        table = (
            data.groupby("group")
            .agg(
                loans=(target, "size"),
                defaults=(target, "sum"),
                observed_default_rate=(target, "mean"),
                mean_predicted_probability=(probability, "mean"),
            )
            .reset_index()
        )
        table["observed_to_predicted"] = (
            table["observed_default_rate"] / table["mean_predicted_probability"]
        )
        table.insert(0, "check", check_name)
        frames.append(table)

    return pd.concat(frames, ignore_index=True)


# Step 6: Save plain result tables and the report figure

def prepare_metrics(metrics):
    table = metrics.copy()
    table["period"] = table["period"].map(PERIOD_NAMES)
    table["model"] = table["model"].map(MODEL_NAMES)
    return table.rename(
        columns={
            "period": "Validation or test period",
            "model": "Model",
            "loans": "Loans",
            "defaults": "Defaults",
            "roc_auc": "ROC AUC",
            "brier_score": "Brier score",
            "mean_probability": "Mean predicted default probability",
            "default_rate": "Observed default rate",
        }
    )


def prepare_calibration(calibration):
    return calibration.rename(
        columns={
            "check": "Check",
            "group": "Group",
            "loans": "Loans",
            "defaults": "Defaults",
            "observed_default_rate": "Observed default rate",
            "mean_predicted_probability": "Mean predicted default probability",
            "observed_to_predicted": "Observed to predicted ratio",
        }
    )


def draw_figure(metrics, calibration):
    test_metrics = metrics.loc[metrics["period"] == "test"].set_index("model")
    order = list(MODEL_NAMES)
    short_names = {
        "public_logistic": "Public logistic",
        "public_pricing_logistic": "Public and pricing logistic",
        "rating_logistic": "Rating logistic",
        "public_pricing_boosting": "Public and pricing boosting",
    }
    risk_groups = calibration.loc[calibration["check"] == "Risk group"]
    countries = calibration.loc[calibration["check"] == "Country"]
    ranking = test_metrics.loc[order].sort_values("roc_auc")

    fig, axes = plt.subplot_mosaic(
        [["ranking", "ranking"], ["risk", "country"]],
        figsize=(8.2, 7.0),
        gridspec_kw={"height_ratios": [0.85, 1.15]},
        layout="constrained",
    )

    ranking_bars = axes["ranking"].barh(
        [short_names[name] for name in ranking.index],
        ranking["roc_auc"],
        color="#4472C4",
    )
    axes["ranking"].bar_label(ranking_bars, fmt="%.3f", padding=4, fontsize=9)
    axes["ranking"].set_title("(a) Model ranking on the test set")
    axes["ranking"].set_xlabel("ROC AUC")
    axes["ranking"].set_xlim(0.60, 0.72)

    upper = 1.05 * max(
        risk_groups["observed_default_rate"].max(),
        risk_groups["mean_predicted_probability"].max(),
    )
    axes["risk"].plot(
        risk_groups["mean_predicted_probability"],
        risk_groups["observed_default_rate"],
        marker="o",
        color="#0072B2",
    )
    axes["risk"].plot([0, upper], [0, upper], "k--", linewidth=1)
    axes["risk"].set_xlabel("Mean predicted probability")
    axes["risk"].set_ylabel("Observed default rate")
    axes["risk"].set_title("(b) Ten probability groups")
    axes["risk"].xaxis.set_major_formatter(PercentFormatter(1.0))
    axes["risk"].yaxis.set_major_formatter(PercentFormatter(1.0))

    positions = np.arange(len(countries))
    width = 0.36
    observed_bars = axes["country"].bar(
        positions - width / 2,
        countries["observed_default_rate"],
        width,
        label="Observed",
        color="0.55",
    )
    predicted_bars = axes["country"].bar(
        positions + width / 2,
        countries["mean_predicted_probability"],
        width,
        label="Predicted",
        color="#0072B2",
    )
    axes["country"].bar_label(
        observed_bars,
        labels=[f"{value:.1%}" for value in countries["observed_default_rate"]],
        padding=2,
        fontsize=8,
    )
    axes["country"].bar_label(
        predicted_bars,
        labels=[
            f"{value:.1%}" for value in countries["mean_predicted_probability"]
        ],
        padding=2,
        fontsize=8,
    )
    axes["country"].set_xticks(positions, countries["group"], rotation=15)
    axes["country"].set_ylabel("Default rate")
    axes["country"].set_title("(c) Country check")
    axes["country"].yaxis.set_major_formatter(PercentFormatter(1.0))
    axes["country"].legend(frameon=False)

    fig.savefig(FIGURE)
    plt.close(fig)


def main():
    loans = create_features(load_data())
    main_loans = loans.loc[loans["loan_group"] == "main"]
    denmark = loans.loc[loans["loan_group"] == "Denmark"]
    train = main_loans.loc[main_loans["split"] == "train"]
    validation = main_loans.loc[main_loans["split"] == "validation"]
    test = main_loans.loc[main_loans["split"] == "test"]

    models, metrics, test_predictions = fit_models(train, validation, test)
    target = "default_within_12_months"
    denmark_predictions = denmark[
        ["loan_id", "loan_issued_at", "country", target]
    ].copy()
    denmark_predictions["public_pricing_logistic"] = models[
        "public_pricing_logistic"
    ].predict_proba(denmark)[:, 1]
    calibration = calibration_table(test_predictions, denmark_predictions)

    with pd.ExcelWriter(RESULTS) as writer:
        prepare_metrics(metrics).to_excel(writer, sheet_name="model_results", index=False)
        prepare_calibration(calibration).to_excel(
            writer, sheet_name="calibration_checks", index=False
        )

    EXTRAS_DIR.mkdir(exist_ok=True)
    draw_figure(metrics, calibration)
    print(prepare_metrics(metrics).to_string(index=False))


if __name__ == "__main__":
    main()
